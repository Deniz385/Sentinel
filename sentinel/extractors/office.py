"""Office formats: DOCX, XLSX/XLSM, PPTX.

In the previous version `.docx` was in the plain-text extension list,
so ZIP binary content was treated as text and nothing was ever found.
Here each format is read through its real structure; without the
library, DOCX falls back to raw XML.

XLSX is opened with ``read_only=True``, so a 100 MB workbook is never
fully loaded into memory.
"""

from __future__ import annotations

import re
import zipfile

from sentinel.extractors import TextChunk

DOCX_EXTENSIONS = frozenset({".docx", ".docm", ".dotx"})
XLSX_EXTENSIONS = frozenset({".xlsx", ".xlsm", ".xltx"})
PPTX_EXTENSIONS = frozenset({".pptx", ".pptm", ".potx"})

#: Upper bound on cells emitted per row (for absurdly wide sheets).
_MAX_CELLS_PER_ROW = 500
_ROWS_PER_CHUNK = 2000


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------
_XML_TAG_RE = re.compile(r"<[^>]+>")
_PARA_END_RE = re.compile(r"</w:p>")


def _docx_via_xml(path: str):
    """Read ``word/document.xml`` directly when python-docx is absent.

    Headers, footers and notes are scanned too; secrets often sit there
    (a test account left in a template header, say).
    """
    targets = ("word/document.xml",)
    with zipfile.ZipFile(path) as archive:
        names = [
            name
            for name in archive.namelist()
            if name in targets
            or re.match(r"word/(header|footer|footnotes|endnotes|comments)\d*\.xml", name)
        ]
        for name in names:
            try:
                raw = archive.read(name).decode("utf-8", errors="replace")
            except Exception:
                continue
            # Turn paragraph ends into newlines, drop the remaining tags.
            raw = _PARA_END_RE.sub("\n", raw)
            text = _XML_TAG_RE.sub("", raw)
            if text.strip():
                yield text


def extract_docx(path: str, config):
    try:
        import docx  # python-docx
    except ImportError:
        docx = None

    if docx is None:
        try:
            for text in _docx_via_xml(path):
                yield TextChunk(text, path, 0, path)
        except (zipfile.BadZipFile, OSError):
            return
        return

    try:
        document = docx.Document(path)
    except Exception:
        # Corrupt or password-protected file: try the XML path.
        try:
            for text in _docx_via_xml(path):
                yield TextChunk(text, path, 0, path)
        except Exception:
            return
        return

    lines = [paragraph.text for paragraph in document.paragraphs if paragraph.text]

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                lines.append(" | ".join(cells))

    for section in document.sections:
        for container in (section.header, section.footer):
            if container is None:
                continue
            for paragraph in container.paragraphs:
                if paragraph.text:
                    lines.append(paragraph.text)

    if lines:
        yield TextChunk("\n".join(lines), path, 0, path)


# ---------------------------------------------------------------------------
# XLSX
# ---------------------------------------------------------------------------
def extract_xlsx(path: str, config):
    try:
        import openpyxl
    except ImportError:
        return

    try:
        workbook = openpyxl.load_workbook(
            path, data_only=True, read_only=True, keep_links=False
        )
    except Exception:
        return

    try:
        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            buffer: list[str] = []
            offset = 0
            for row in sheet.iter_rows(values_only=True):
                cells = [
                    str(cell) for cell in row[:_MAX_CELLS_PER_ROW] if cell is not None
                ]
                buffer.append(" | ".join(cells))
                if len(buffer) >= _ROWS_PER_CHUNK:
                    yield TextChunk(
                        "\n".join(buffer), f"{path}::{sheet_name}", offset, path
                    )
                    offset += len(buffer)
                    buffer = []
            if any(line.strip() for line in buffer):
                yield TextChunk(
                    "\n".join(buffer), f"{path}::{sheet_name}", offset, path
                )
    finally:
        try:
            workbook.close()
        except Exception:
            pass


# ---------------------------------------------------------------------------
# PPTX
# ---------------------------------------------------------------------------
def extract_pptx(path: str, config):
    try:
        from pptx import Presentation
    except ImportError:
        return

    try:
        presentation = Presentation(path)
    except Exception:
        return

    for index, slide in enumerate(presentation.slides, 1):
        lines: list[str] = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                for paragraph in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in paragraph.runs)
                    if text.strip():
                        lines.append(text)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if cells:
                        lines.append(" | ".join(cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text
            if notes and notes.strip():
                lines.append(notes)
        if lines:
            yield TextChunk("\n".join(lines), f"{path}::slide-{index}", 0, path)
